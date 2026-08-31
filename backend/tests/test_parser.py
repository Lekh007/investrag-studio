import sys
from pathlib import Path
from types import SimpleNamespace

from investrag.parser import extract_email_attachments, parse_document


def test_json_provenance(tmp_path: Path) -> None:
    path = tmp_path / "payload.json"
    path.write_text('{"company":{"name":"Acme","metrics":[{"revenue":42}]}}', encoding="utf-8")
    parsed = parse_document(path)
    assert parsed.parser == "native-json"
    assert any(element.json_path == "$.company.metrics[0].revenue" for element in parsed.elements)


def test_markdown_headings_and_lists(tmp_path: Path) -> None:
    path = tmp_path / "notes.md"
    path.write_text("# Revenue\n\n- Growth was 12%.\n", encoding="utf-8")
    parsed = parse_document(path)
    assert [element.element_type for element in parsed.elements] == ["heading", "list"]


def test_csv_and_html_keep_structural_elements(tmp_path: Path) -> None:
    csv_path = tmp_path / "metrics.csv"
    csv_path.write_text("metric,value\nmargin,42%\n", encoding="utf-8")
    html_path = tmp_path / "brief.html"
    html_path.write_text("<html><script>ignore()</script><h1>Outlook</h1><p>Stable demand.</p></html>", encoding="utf-8")

    csv_elements = parse_document(csv_path).elements
    html_elements = parse_document(html_path).elements
    assert csv_elements[0].cell_range == "header"
    assert csv_elements[1].cell_range == "row:2"
    assert [element.text for element in html_elements] == ["Outlook", "Stable demand."]


def test_pdf_and_xlsx_include_location_provenance(tmp_path: Path) -> None:
    import openpyxl
    import pymupdf as fitz

    pdf_path = tmp_path / "report.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Revenue 42")
    document.save(pdf_path)
    document.close()

    workbook_path = tmp_path / "metrics.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Summary"
    sheet["A1"] = "Revenue"
    sheet["B1"] = 42
    workbook.save(workbook_path)

    pdf_elements = parse_document(pdf_path).elements
    xlsx_elements = parse_document(workbook_path).elements
    assert any(element.page == 1 and element.bbox for element in pdf_elements)
    assert any(element.sheet == "Summary" and element.cell_range == "A1:B1" for element in xlsx_elements)


def test_content_signature_routes_a_renamed_pdf(tmp_path: Path) -> None:
    import pymupdf

    disguised = tmp_path / "report.txt"
    document = pymupdf.open()
    document.new_page().insert_text((72, 72), "Revenue 42")
    document.save(disguised)
    document.close()

    parsed = parse_document(disguised)
    assert parsed.media_type == "application/pdf"
    # The parser that wins is an adaptive-routing detail (docling when the optional
    # `parsers` extra is installed, pymupdf-native otherwise); content-signature
    # detection is proven by media_type, not by which PDF parser handled it.
    assert parsed.parser in {"docling", "pymupdf-native"}


def test_generated_office_demo_files_are_valid_packages(tmp_path: Path) -> None:
    from docx import Document
    from pptx import Presentation
    from scripts.create_demo_corpus import create_corpus

    create_corpus(tmp_path)
    assert "Memo: risk is stable." in "\n".join(
        paragraph.text for paragraph in Document(tmp_path / "memo.docx").paragraphs
    )
    presentation = Presentation(tmp_path / "update.pptx")
    assert any("margin expanded" in shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text"))


def test_msg_uses_dedicated_parser_and_attachment_path(tmp_path: Path, monkeypatch) -> None:
    class FakeAttachment:
        data = b"attachment body"

        def getFilename(self) -> str:
            return "note.txt"

    class FakeMessage:
        subject = "Investment update"
        sender = "analyst@example.com"
        to = "pm@example.com"
        date = "2026-08-29"
        body = "Revenue increased."
        htmlBody = None
        attachments = [FakeAttachment()]

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

    fake_module = SimpleNamespace(openMsg=lambda *_args, **_kwargs: FakeMessage())
    monkeypatch.setitem(sys.modules, "extract_msg", fake_module)
    path = tmp_path / "note.msg"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1fake")

    parsed = parse_document(path)
    assert parsed.parser == "extract-msg"
    assert any(element.text == "Revenue increased." for element in parsed.elements)
    assert extract_email_attachments(path) == [("note.txt", b"attachment body")]
